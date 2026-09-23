#!/usr/bin/env python3
"""Converte a planilha socioassistencial em dados consumidos pelo mapa.

Uso, a partir da pasta raiz do projeto:
    python scripts/gerar_dados.py
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "data" / "mapeamento_smas.xlsx"
OUTPUT = ROOT / "data" / "equipamentos.js"

# Aliases de tipo encontrados na planilha. Os nomes são normalizados abaixo.
TIPOS_NORMALIZADOS = {
    "CADRIO": "CADRIO",
    "CONTUTELAR": "CONSELHO_TUTELAR",
    "CENTROPOP": "CENTRO_POP",
}

CAMPOS_OBRIGATORIOS = {
    "id", "cas", "tipo", "nome", "cas_nome", "endereco", "bairro",
    "cep", "email", "responsavel", "abrangencia", "latitude_aproximada",
    "longitude_aproximada", "precisao_geocodificacao",
}


def texto(valor: object) -> str:
    """Converte células vazias/numéricas em texto sem expor None no mapa."""
    if valor is None:
        return ""
    return str(valor).strip()


def ler_planilha() -> list[dict[str, object]]:
    """Lê a primeira aba e valida cabeçalhos e coordenadas antes da conversão."""
    if not INPUT.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {INPUT}")

    folha = load_workbook(INPUT, data_only=True, read_only=True).active
    linhas = folha.iter_rows(values_only=True)
    cabecalhos = [texto(c) for c in next(linhas)]
    ausentes = sorted(CAMPOS_OBRIGATORIOS - set(cabecalhos))
    if ausentes:
        raise ValueError("Colunas obrigatórias ausentes: " + ", ".join(ausentes))

    equipamentos = []
    for numero_linha, valores in enumerate(linhas, start=2):
        linha = dict(zip(cabecalhos, valores))
        if not texto(linha.get("nome")):
            continue

        try:
            latitude = float(texto(linha.get("latitude_aproximada")).replace(",", "."))
            longitude = float(texto(linha.get("longitude_aproximada")).replace(",", "."))
        except (TypeError, ValueError) as erro:
            raise ValueError(
                f"Latitude/longitude inválidas na linha {numero_linha}: "
                f"{linha.get('nome')}"
            ) from erro

        if not (-90 <= latitude <= 90 and -180 <= longitude <= 180):
            raise ValueError(f"Coordenadas fora do intervalo na linha {numero_linha}.")

        tipo_planilha = texto(linha.get("tipo")).upper()
        tipo = TIPOS_NORMALIZADOS.get(tipo_planilha, tipo_planilha)
        if not tipo:
            raise ValueError(f"Tipo ausente na linha {numero_linha}.")

        propriedades = {
            "id": texto(linha.get("id")),
            "cas": texto(linha.get("cas")),
            "tipo": tipo,
            "nome": texto(linha.get("nome")),
            "cas_nome": texto(linha.get("cas_nome")),
            "endereco": texto(linha.get("endereco")),
            "bairro": texto(linha.get("bairro")),
            "cep": texto(linha.get("cep")),
            "email": texto(linha.get("email")),
            "responsavel": texto(linha.get("responsavel")),
            "abrangencia": texto(linha.get("abrangencia")),
            "precisao": texto(linha.get("precisao_geocodificacao")),
        }
        equipamentos.append({
            "type": "Feature",
            "geometry": {"type": "Point", "coordinates": [longitude, latitude]},
            "properties": propriedades,
        })

    if not equipamentos:
        raise ValueError("A planilha não contém equipamentos com nome preenchido.")
    return equipamentos


def main() -> None:
    equipamentos = ler_planilha()
    conteudo = "/* Gerado por scripts/gerar_dados.py. Edite a planilha, não este arquivo. */\n"
    conteudo += "window.SMAS_FEATURES = "
    conteudo += json.dumps(equipamentos, ensure_ascii=False, separators=(",", ":"))
    conteudo += ";\n"
    OUTPUT.write_text(conteudo, encoding="utf-8")
    print(f"Dados gerados: {len(equipamentos)} equipamentos → {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
